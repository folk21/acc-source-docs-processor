"""Parser for 1C-style XLSX account-card expense statements."""

from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook

from .models import StatementEntry


_PERSON_TOKEN = re.compile(r"^[A-Za-zА-Яа-яЁё][A-Za-zА-Яа-яЁё.'’-]*$")


def _normalize_header(value: object) -> str:
    """Normalize one workbook header value for structural matching."""
    return re.sub(r"\s+", " ", str(value or "").casefold().replace("ё", "е")).strip()


def _parse_date(value: object) -> date | None:
    """Parse common Excel or text date values."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str):
        return None
    text = value.strip()
    for pattern in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def _parse_decimal(value: object) -> Decimal | None:
    """Convert one numeric XLSX value to a two-decimal Decimal."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        result = value
    elif isinstance(value, (int, float)):
        result = Decimal(str(value))
    elif isinstance(value, str):
        normalized = value.replace("\u00a0", "").replace(" ", "").replace(",", ".")
        if not normalized:
            return None
        try:
            result = Decimal(normalized)
        except InvalidOperation:
            return None
    else:
        return None
    return result.quantize(Decimal("0.01"))


def _find_header_row(worksheet) -> tuple[int, dict[str, int]]:
    """Find the account-card group header and its important columns."""
    for row_number in range(1, min(worksheet.max_row, 40) + 1):
        values = {
            column: _normalize_header(worksheet.cell(row_number, column).value)
            for column in range(1, worksheet.max_column + 1)
        }
        period = next((column for column, value in values.items() if value.startswith("период")), None)
        document = next((column for column, value in values.items() if value.startswith("документ")), None)
        debit = next((column for column, value in values.items() if value == "дебет"), None)
        credit = next((column for column, value in values.items() if value == "кредит"), None)
        balance = next((column for column, value in values.items() if "сальдо" in value), None)
        if period and document and credit:
            return row_number, {
                "period": period,
                "document": document,
                "debit": debit or credit,
                "credit": credit,
                "balance": balance or worksheet.max_column + 1,
            }
    raise ValueError("Cannot find Period/Document/Credit headers in the statement workbook")


def _data_rows(worksheet, header_row: int, period_column: int) -> list[int]:
    """Return rows that represent dated statement transactions."""
    return [
        row_number
        for row_number in range(header_row + 1, worksheet.max_row + 1)
        if _parse_date(worksheet.cell(row_number, period_column).value) is not None
    ]


def _find_credit_amount_column(
    worksheet,
    rows: list[int],
    credit_column: int,
    balance_column: int,
) -> int:
    """Select the numeric credit subcolumn with transaction-like value diversity."""
    candidates: list[tuple[tuple[int, int, int], int]] = []
    for column in range(credit_column, min(balance_column, worksheet.max_column + 1)):
        values = [
            value
            for row in rows
            if (value := _parse_decimal(worksheet.cell(row, column).value)) is not None
            and value != 0
        ]
        if not values:
            continue
        unique_count = len(set(values))
        candidates.append(((unique_count, len(values), int(column > credit_column)), column))
    if not candidates:
        raise ValueError("Cannot find a numeric credit amount column in the statement workbook")
    return max(candidates)[1]


def _looks_like_person(value: object) -> bool:
    """Return True for compact two-to-five-token person-like workbook values."""
    if not isinstance(value, str) or "\n" in value or len(value) > 80:
        return False
    tokens = [token for token in value.strip().split() if token]
    if not 2 <= len(tokens) <= 5:
        return False
    return all(_PERSON_TOKEN.fullmatch(token) for token in tokens)


def _find_person_column(
    worksheet,
    rows: list[int],
    document_column: int,
    debit_column: int,
) -> int | None:
    """Infer the employee/person analytics column between Document and Debit."""
    candidates: list[tuple[tuple[int, float], int]] = []
    for column in range(document_column + 1, debit_column):
        values = [worksheet.cell(row, column).value for row in rows]
        matches = [str(value).strip() for value in values if _looks_like_person(value)]
        if not matches:
            continue
        average_length = mean(len(value) for value in matches)
        candidates.append(((len(matches), -average_length), column))
    return max(candidates)[1] if candidates else None


def parse_statement(path: Path) -> list[StatementEntry]:
    """Parse expense positions from a 1C-style XLSX account-card statement."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        header_row, columns = _find_header_row(worksheet)
        rows = _data_rows(worksheet, header_row, columns["period"])
        if not rows:
            raise ValueError("Statement workbook contains no dated transaction rows")
        amount_column = _find_credit_amount_column(
            worksheet,
            rows,
            columns["credit"],
            columns["balance"],
        )
        person_column = _find_person_column(
            worksheet,
            rows,
            columns["document"],
            columns["debit"],
        )

        entries: list[StatementEntry] = []
        for row_number in rows:
            transaction_date = _parse_date(
                worksheet.cell(row_number, columns["period"]).value
            )
            amount = _parse_decimal(worksheet.cell(row_number, amount_column).value)
            if transaction_date is None or amount is None or amount <= 0:
                continue
            person_name = None
            if person_column is not None:
                raw_person = worksheet.cell(row_number, person_column).value
                if raw_person not in {None, ""}:
                    person_name = str(raw_person).strip()
            raw_description = worksheet.cell(row_number, columns["document"]).value
            entries.append(
                StatementEntry(
                    source_row=row_number,
                    transaction_date=transaction_date,
                    person_name=person_name,
                    amount=amount,
                    description=(
                        str(raw_description).strip()
                        if raw_description not in {None, ""}
                        else None
                    ),
                )
            )
        if not entries:
            raise ValueError("Statement workbook contains no positive credit transactions")
        return entries
    finally:
        workbook.close()

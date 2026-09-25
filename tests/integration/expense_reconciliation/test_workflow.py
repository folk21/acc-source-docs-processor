"""Synthetic end-to-end expense reconciliation coverage."""

from datetime import date
from decimal import Decimal

import fitz
from openpyxl import Workbook, load_workbook

from source_docs_processor.features.expense_reconciliation import reconcile_expenses


def _write_statement(path) -> None:
    """Write a synthetic 1C-style statement with one intentionally missing receipt."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Карточка счета 51"])
    worksheet.append(
        [
            "Период",
            "Документ",
            "Аналитика",
            "Дебет",
            None,
            None,
            "Кредит",
            None,
            None,
            "Текущее сальдо",
        ]
    )
    worksheet.append([None, None, None, "Счет", None, None, "Счет"])
    worksheet.append([date(2026, 7, 2), "Payment 1", "Учебный Сотрудник", "71.01", None, None, "51", 1000])
    worksheet.append([date(2026, 7, 2), "Payment 2", "Учебный Сотрудник", "71.01", None, None, "51", 25])
    worksheet.append([date(2026, 7, 3), "Payment 3", "Учебный Сотрудник", "71.01", None, None, "51", 500])
    workbook.save(path)


def _write_text_pdf(path, text: str) -> None:
    """Write a deterministic native-text PDF that never requires real OCR."""
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    document.save(path)
    document.close()


def test_reconciliation_pipeline_writes_matches_documents_and_summary(tmp_path) -> None:
    """Verify source folder plus statement produces the requested XLSX report.

    Protected risk: the new operation must reconcile across both inputs, support
    a one-document-to-two-position match, retain one missing statement row, and
    avoid counting the combined document amount twice in aggregate totals.
    """
    source = tmp_path / "documents"
    source.mkdir()
    _write_text_pdf(
        source / "ticket.pdf",
        "E-TICKET\nPassenger name\nSTUDENT/USER MR\nIssue date 02.07.2026\nTOTAL 1025.00",
    )
    statement = tmp_path / "payments.xlsx"
    _write_statement(statement)
    output = tmp_path / "output"

    summary = reconcile_expenses(source, statement, output)

    assert summary.document_count == 1
    assert summary.statement_position_count == 3
    assert summary.matched_statement_position_count == 2
    assert summary.unmatched_statement_position_count == 1
    assert summary.statement_total == Decimal("1525.00")
    assert summary.document_total == Decimal("1025.00")
    assert summary.difference == Decimal("500.00")
    assert summary.unmatched_statement_total == Decimal("500.00")
    assert summary.workbook_path == output / "expense_reconciliation.xlsx"

    workbook = load_workbook(summary.workbook_path, data_only=True)
    reconciliation = workbook["Reconciliation"]
    documents = workbook["Documents"]
    assert [reconciliation.cell(row, 4).value for row in (2, 3, 4)] == [
        "Найден",
        "Найден",
        "Не найден",
    ]
    assert [reconciliation.cell(row, 5).value for row in (2, 3)] == [1, 2]
    assert [reconciliation.cell(row, 7).value for row in (2, 3)] == [1025, 1025]
    assert documents.cell(2, 1).value == "ticket.pdf"
    assert documents.cell(2, 6).value == "Да"
    assert documents.cell(2, 7).value == 2
    workbook.close()

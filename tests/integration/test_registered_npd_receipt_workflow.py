import csv
from pathlib import Path

import cv2
import numpy as np
from openpyxl import load_workbook

from source_docs_processor.cli import process_folder
from source_docs_processor.document_processor import BaseDocumentProcessor
from source_docs_processor.models import ExtractedDocument


class FakeNpdReceiptProcessor(BaseDocumentProcessor):
    """Return deterministic receipt results without calling Tesseract."""

    document_type = "npd_receipts"
    display_name = "Fake NPD receipt processor"

    def analyze_image_orientations(
        self,
        image_path: Path,
        image: np.ndarray,
        **_kwargs,
    ):
        """Recognize one synthetic file and reject all other images."""
        if image_path.name == "incoming.JPG":
            return (
                ExtractedDocument(
                    source_path=image_path,
                    document_type=self.document_type,
                    is_recognized=True,
                    document_date="2026-07-15",
                    document_number="receipt-42",
                    total_amount="1250.00",
                    recipient_name="ООО Учебный Покупатель",
                    recipient_inn="0000000000",
                    confidence=98,
                ),
                image,
            )
        return (
            ExtractedDocument(
                source_path=image_path,
                document_type=self.document_type,
            ),
            image,
        )


def _write_test_image(path: Path) -> None:
    """Create a tiny valid image for receipt workflow integration testing."""
    path.parent.mkdir(parents=True, exist_ok=True)
    image = np.full((12, 12, 3), 255, dtype=np.uint8)
    assert cv2.imwrite(str(path), image)


def test_registered_receipt_workflow_copies_only_receipts_and_writes_registries(tmp_path):
    """Verify the new document type owns its copy and short-registry behavior.

    Protected risk: the receipt workflow must not copy unrelated images or
    inherit the detailed UPD registry and continuation-page behavior.
    """
    source_dir = tmp_path / "source"
    output_dir = tmp_path / "output"
    _write_test_image(source_dir / "nested" / "incoming.JPG")
    _write_test_image(source_dir / "nested" / "other.png")

    found, all_documents = process_folder(
        source_dir=source_dir,
        output_dir=output_dir,
        lang="rus+eng",
        document_type="npd_receipts",
        document_processor=FakeNpdReceiptProcessor(),
    )

    target_dir = output_dir / "чеки_к_загрузке"
    copied_name = "2026-07-15_1250-00_receipt-42.JPG"
    assert len(found) == 1
    assert len(all_documents) == 2
    assert (target_dir / copied_name).exists()
    assert not (target_dir / "other.png").exists()

    csv_path = target_dir / "чеки_к_загрузке.csv"
    with csv_path.open("r", encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file, delimiter=";"))

    expected_columns = [
        "document_date",
        "total_amount",
        "recipient_name",
        "document_number",
        "copied_file_name",
        "source_file_name",
        "recipient_inn",
    ]
    assert list(rows[0]) == expected_columns
    assert rows[0]["copied_file_name"] == copied_name
    assert rows[0]["source_file_name"] == "incoming.JPG"

    workbook = load_workbook(target_dir / "чеки_к_загрузке.xlsx")
    worksheet = workbook["Receipts"]
    assert [cell.value for cell in worksheet[1]] == expected_columns
    assert worksheet["A2"].value == "2026-07-15"
    assert worksheet["B2"].value == 1250.0
    assert worksheet["E2"].value == copied_name
